from flask import Flask, render_template, request, send_file
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import os
from openai import OpenAI
from dotenv import load_dotenv
from datetime import datetime

load_dotenv()
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

app = Flask(__name__)

grade_styles = {
    "초등 저학년": {
        "조건": lambda g: any(x in g for x in ["초1", "초2"]),
        "스타일": "따뜻하고 귀여운 말투로, 아이의 즐거운 수업 참여와 성장을 중심으로 학부모에게 전달해주세요."
    },
    "초등 고학년": {
        "조건": lambda g: any(x in g for x in ["초3", "초4", "초5", "초6"]),
        "스타일": "친근하지만 논리적인 어조로, 수업 참여도와 성취 과정을 중심으로 정리해주세요."
    },
    "중학생": {
        "조건": lambda g: "중" in g,
        "스타일": "차분하고 구체적인 말투로, 자기주도성과 협업 능력 등을 포함해 작성해주세요."
    },
    "고등학생": {
        "조건": lambda g: "고" in g,
        "스타일": "진로 및 성취 중심으로, 목표 설정과 실제 수행 과정의 연결이 느껴지도록 전문적인 어조로 작성해주세요."
    },
    "기본": {
        "조건": lambda g: True,
        "스타일": "학부모님께 공감 가는 친절한 말투로, 학생의 수업 참여 및 향후 일정에 대해 전달해주세요."
    }
}


def get_style_by_grade(grade):
    for style_name, info in grade_styles.items():
        if info["조건"](grade):
            return style_name, info["스타일"]
    return "기본", grade_styles["기본"]["스타일"]


evaluation_phrases = {
    "참여도": {"전혀": "수업 참여가 거의 없었어요.", "조금": "가끔 참여하는 모습을 보였어요.", "보통": "일반적인 수준으로 참여했어요.", "좋음": "적극적으로 참여했어요.", "매우 좋음": "항상 열정적으로 참여했어요. 😊"},
    "이해도": {"전혀": "내용을 이해하는 데 어려움이 있었어요.", "조금": "기초 개념을 부분적으로 이해했어요.", "보통": "기본 개념을 이해했어요.", "좋음": "개념을 잘 이해했어요.", "매우 좋음": "내용을 빠르게 파악하고 정확히 이해했어요. 👍"},
    "컴퓨터 사용능력": {"전혀": "기초적인 컴퓨터 조작에 어려움이 있었어요.", "조금": "간단한 조작이 가능했어요.", "보통": "일반적인 사용이 가능했어요.", "좋음": "능숙하게 활용했어요.", "매우 좋음": "활용 능력이 매우 뛰어났어요. 💻"},
    "디버깅": {"전혀": "문제 해결을 시도하지 않았어요.", "조금": "간단한 오류를 해결하려 했어요.", "보통": "일반적인 오류를 해결했어요.", "좋음": "오류를 잘 탐지하고 해결했어요.", "매우 좋음": "복잡한 오류도 논리적으로 해결했어요. 🛠️"},
    "응용력": {"전혀": "배운 내용을 적용하지 못했어요.", "조금": "간단한 문제에 적용 가능했어요.", "보통": "기초적인 응용은 가능했어요.", "좋음": "다양한 문제에 응용했어요.", "매우 좋음": "새로운 문제에 창의적으로 응용했어요. 🌟"},
    "의사소통능력": {"전혀": "소통에 어려움이 있었어요.", "조금": "간단히 표현했어요.", "보통": "기본적인 소통이 가능했어요.", "좋음": "의견을 잘 전달했어요.", "매우 좋음": "적극적으로 소통하며 협력했어요. 🤝"}
}

def convert_evaluation_to_text(evaluations):
    return " ".join([
        evaluation_phrases.get(k, {}).get(v, "")
        for k, v in evaluations.items() if v
    ])

class NoticeWord:
    def __init__(self):
        self.doc = Document()
        self.doc.add_heading('📄 가정통신문', 0)

    def create_notice(self, data):
        for key, value in data.items():
            if value:
                para = self.doc.add_paragraph()
                run = para.add_run(f"{key}: {value}" if key != "제목" and key != "본문" else f"{value}")
                run.font.size = Pt(12)
                run.font.name = '맑은 고딕'
                r = run._element
                r.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    def save(self, filepath):
        self.doc.save(filepath)


def generate_ai_text(data, role="본문"):

    style_name, style_instruction = get_style_by_grade(data.get("학년", ""))

    system_msg = (
        f"너는 학생의 학부모에게 전달하는 가정통신문의 '{role}' 부분을 작성하는 AI야.\n\n"
        f"학생의 학년은 '{data.get('학년')}', 스타일 분류는 '{style_name}'야.\n"
        f"{style_instruction}\n\n"
        f"학생의 수업 정보:\n"
        f"- 이름: {data.get('이름')}\n"
        f"- 학년: {data.get('학년')}\n"
        f"- 수업: {data.get('수업')}\n"
        f"- 수업 중 주요 내용 또는 메모:\n{data.get('본문', '')}\n\n"
        f"- 본문에 들어가야 할 평가 내용: {data.get('평가', '')}\n\n"
        f"- 향후 일정 안내: {data.get('향후 일정 안내', '')}\n\n"
        "작성 시 주의사항:\n"
        "- 처음 시작시 계절의 내용과 함께, 인삿말 건네고 시작. 학부모님께 긍정적이고 명확한 어조로 안내해줘.\n"
        "- 제목, 평가(\"평가:~ \" 이런 내용 절대 절대 표시하지 말고, 본문 안에 자연스럽게 포함시켜서 표현해줘.\n"
        "- 필요시 이모지도 사용하고, 마지막엔 수원디랩코딩학원 드림 문장으로 끝내줘.\n"
        "- 감사인사 끝난 후, 향후 일정을 표시할 때, 본문의 마지막 부분에, 📅 <향후 일정 안내> 이렇게 사용하고 다음 줄부터 줄바꿈 없이 빠뜨리지 말고 전부 가독성 있게 나열해줘.\n"
    )

    user_msg = f"{data.get('이름')} 학생은 {data.get('수업')} 수업에 참여했어요."
    
    response = client.chat.completions.create(
        model="gpt-4",
        messages=[
            {"role": "system", "content": system_msg},
            {"role": "user", "content": user_msg}
        ],
        temperature=0.7,
        max_tokens=600
    )

    return response.choices[0].message.content.strip()

@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        year = request.form.get("year")
        season = request.form.get("season")
        title_text = request.form.get("title")
        full_title = f"{year} {season} {title_text}"

        name = request.form.get("name")
        grade = request.form.get("grade")
        class_name = request.form.get("class")
        date = request.form.get("date")
        body_text = request.form.get("body")
        event = request.form.get("event")

        evaluations = {
            "참여도": request.form.get("eval_participation"),
            "이해도": request.form.get("eval_understanding"),
            "컴퓨터 사용능력": request.form.get("eval_computer"),
            "디버깅": request.form.get("eval_debug"),
            "응용력": request.form.get("eval_application"),
            "의사소통능력": request.form.get("eval_communication")
        }

        converted_eval = convert_evaluation_to_text(evaluations)
        event_info = request.form.get("event", "")
        body_text = request.form.get("body")
        ai_generate_body = request.form.get("generate_body")

        data = {
                "이름": name,
                "학년": grade,
                "수업": class_name,
                "날짜": date,
                "제목": full_title,
                "평가": evaluations,
                "본문": body_text,
                "향후 일정 안내": event_info
        }

        if ai_generate_body == '1':
            body_text = generate_ai_text(data)
            eval_text = f'📝 <학생 수업 요약>\n{converted_eval}'
            data["본문"] = (body_text.strip() if body_text else "") + "\n\n" + eval_text
    
        del data["날짜"]
        del data["평가"]  # Remove the evaluations dict to avoid duplication in the document
        del data["향후 일정 안내"]  # Remove the event info to avoid duplication in the document
        
        word = NoticeWord()
        word.create_notice(data)

        os.makedirs("output", exist_ok=True)
        output_path = os.path.join("output", f"{name} {full_title}.docx")
        word.save(output_path)

        return send_file(output_path, as_attachment=True)

    return render_template("index.html")

if __name__ == "__main__":
    app.run(debug=True)
